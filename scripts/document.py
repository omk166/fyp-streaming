"""Create a shareable DOCX from real run summaries and screenshots."""
import csv, json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

root=Path('artifacts')
rows=list(csv.DictReader((root/'results.csv').open(encoding='utf-8')))
doc=Document()
doc.add_heading('FYP Streaming Prototype: Initial Local Evaluation',0)
doc.add_paragraph('Measured synthetic-media experiments. This report evaluates the research prototype, not a completed conferencing application.')
doc.add_heading('Environment and method',1)
env=json.loads((root/'environment.json').read_text(encoding='utf-8-sig'))
for key,value in env.items(): doc.add_paragraph(f'{key}: {value}')
doc.add_paragraph('All browser participants and the mediasoup SFU run on the same Windows laptop. Each participant publishes an animated 640 x 360 canvas and a quiet 440 Hz synthetic tone. Browser sessions are isolated. Telemetry is sampled about once per second. Runs include join setup and functional checks before a timed observation period; the tables use all saved samples, including startup. These short trials provide functional evidence and preliminary telemetry, not statistically validated capacity or codec-quality conclusions.')
doc.add_heading('Results',1)
table=doc.add_table(rows=1, cols=7); table.style='Light Shading Accent 1'
headers=['Mode','Peers','Client','Samples','Upload kb/s','ICE RTT p95 ms','Video FPS']
for cell,text in zip(table.rows[0].cells,headers): cell.text=text
def fmt(value): return f'{float(value):.2f}' if value else 'N/A'
for row in rows:
    vals=[row['mode'],row['participants'],row['peer'],row['sample_count'],fmt(row['send_rtp_payload_kbps_mean']),fmt(row['ice_rtt_ms_p95']),fmt(row['inbound_video_fps_mean'])]
    for cell,text in zip(table.add_row().cells,vals): cell.text=text
doc.add_paragraph('Upload is the sum of outbound RTP payload counters differentiated over time; transport headers and network overhead are excluded. FPS averages available inbound video reports. RTT is ICE candidate-pair RTT and does not measure end-to-end audio delay. Per-peer results from a shared laptop are correlated.')
doc.add_heading('Functional checks and run outcomes',1)
for path in sorted((root/'raw').glob('*/manifest.json')):
    run=json.loads(path.read_text())
    doc.add_paragraph(f"{run['runId']}: {'PASS' if run['passed'] else 'FAIL'}; {run['measurementSeconds']} s observation; {run['browserChannel']} {run['browser']}; source {run['gitCommit']}")
    for check in run['checks']: doc.add_paragraph(check['name'],style='List Bullet')
    if run['errors']: doc.add_paragraph('Errors are preserved in the local raw manifest.')
doc.add_heading('Measured figures and screenshots',1)
if (root/'local-measurements.png').exists(): doc.add_picture(str(root/'local-measurements.png'),width=Inches(6))
for image in sorted((root/'screenshots').glob('*.png')):
    doc.add_paragraph(image.stem)
    doc.add_picture(str(image),width=Inches(6))
doc.add_heading('Limitations and pending experiments',1)
for text in [
    'No real microphone, speaker, camera, acoustic echo or mouth-to-ear test has been performed. The original-audio target below 200 ms remains unverified.',
    'No bandwidth caps, packet loss, jitter, WAN delay, NAT or TURN path has been exercised in these local trials.',
    'Multiple browser sessions on one laptop do not establish eight-device meeting capacity or cross-browser compatibility.',
    'No aligned reference/received video recording, VMAF, PSNR, SSIM or human MOS was measured. Synthetic-canvas bitrate cannot establish codec superiority.',
    'Screen sharing, host permissions, whiteboard, recorder and AI branches are not implemented in this prototype and are not tested.',
    'The experiment matrix retains pending work. Further measurements require repeated trials, controlled media, actual network shaping and separate endpoint devices.'
]: doc.add_paragraph(text,style='List Bullet')
doc.add_heading('Evidence and reproduction',1)
doc.add_paragraph('See README.md for startup commands, scripts/smoke.mjs for the automated procedure, scripts/report.py for calculations, artifacts/results.csv for numeric data, and artifacts/screenshots for visual evidence. Raw telemetry and detailed manifests are retained locally under artifacts/raw and excluded from the public repository pending review of connection metadata.')
for section in doc.sections:
    section.header.paragraphs[0].text='FYP Streaming Lab | Preliminary evaluation'
doc.styles['Normal'].font.size=Pt(10)
doc.save(root/'FYP_Streaming_Initial_Evaluation.docx')
print('Saved DOCX report from measured artifacts')
